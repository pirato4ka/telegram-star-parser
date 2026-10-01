"""Тесты сценария `main.run` без реального подключения к Telegram.

Проверяется оркестрация: конфиг -> «авторизация» -> канал -> парсинг -> Excel,
а также корректное завершение по Ctrl+C.
"""

from __future__ import annotations

import argparse
import asyncio
import builtins
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import load_workbook  # noqa: E402
from telethon.tl import types  # noqa: E402

import main as main_module  # noqa: E402
from models import COLUMNS  # noqa: E402
from test_offline import (  # noqa: E402
    FakeClient, make_channel, make_message, make_paid_reactions, make_user,
)

CONFIG_TEMPLATE = """
[API]
ID_API=1
HASH_API=abc
PHONE=+79990000000
CLOUD_PASSWORD=

[DELAY]
MESSAGES_INTERVAL_MIN=0
MESSAGES_INTERVAL_MAX=0

[OUTPUT]
DIR=output
"""


class FlowClient(FakeClient):
    """Клиент с методами, которые нужны сценарию main."""

    async def get_entity(self, peer):
        if isinstance(peer, str):  # username канала из --channel
            return make_channel(1, title="Test Channel", username=peer.lstrip("@"))
        return await super().get_entity(peer)

    async def get_messages(self, entity, limit=1, **kwargs):
        return self.messages[:limit]

    async def disconnect(self):
        return None


class InterruptingClient(FlowClient):
    """Клиент, имитирующий Ctrl+C после первого сообщения."""

    interrupt = KeyboardInterrupt

    async def iter_messages(self, entity, limit=None, offset_id=0, **kwargs):
        for message in self.messages:
            yield message
            raise self.interrupt


class QueueClient(FlowClient):
    """Клиент очереди каналов: у каждого канала свои сообщения.

    `channels` — словарь `username -> (сущность канала, список сообщений)`.
    Канала, которого нет в словаре, «не существует» (как приватный/удалённый).
    """

    def __init__(self, channels=None, entities=None, interrupt_on=None):
        super().__init__(messages=[], entities=entities)
        self.channels = dict(channels or {})
        self.interrupt_on = interrupt_on   # username, на котором «нажимают» Ctrl+C
        self.parsed_order = []

    async def get_entity(self, peer):
        if isinstance(peer, str):
            username = peer.lstrip("@")
            entry = self.channels.get(username)
            if entry is None:
                raise ValueError(f"Could not find the input entity for {peer}")
            return entry[0]
        return await super().get_entity(peer)

    async def get_messages(self, entity, limit=1, **kwargs):
        messages = self._messages_of(entity)
        return messages[:limit]

    async def iter_messages(self, entity, limit=None, offset_id=0, **kwargs):
        username = getattr(entity, "username", "")
        self.parsed_order.append(username)
        if username == self.interrupt_on:
            raise self._interrupt()
        selected = self._messages_of(entity)
        if offset_id:
            selected = [m for m in selected if int(m.id) < offset_id]
        for message in selected[:limit] if limit else selected:
            yield message

    def _messages_of(self, entity):
        return list(self.channels.get(getattr(entity, "username", ""), (None, []))[1])

    @staticmethod
    def _interrupt():
        return KeyboardInterrupt()


class QueueCancellingClient(QueueClient):
    """То же, но настоящая отмена задачи — так Ctrl+C выглядит в asyncio."""

    @staticmethod
    def _interrupt():
        return asyncio.CancelledError()


class FailingChannelClient(QueueClient):
    """Каналы из `fail_on` «падают» посреди парсинга: очередь должна продолжиться."""

    def __init__(self, *args, fail_on=(), **kwargs):
        super().__init__(*args, **kwargs)
        self.fail_on = set(fail_on)

    async def iter_messages(self, entity, limit=None, offset_id=0, **kwargs):
        username = getattr(entity, "username", "")
        if username in self.fail_on:
            self.parsed_order.append(username)
            raise RuntimeError("сбой при чтении истории")
        async for message in super().iter_messages(entity, limit=limit,
                                                   offset_id=offset_id, **kwargs):
            yield message


class CancellingClient(InterruptingClient):
    """Настоящий Ctrl+C: asyncio.run не бросает KeyboardInterrupt, а отменяет задачу."""

    interrupt = asyncio.CancelledError


def make_args(**overrides) -> argparse.Namespace:
    defaults = dict(
        config=None, channel="@testchannel", count=3, output_dir=None,
        warmup=False, create_empty_file=False, log_level="ERROR",
    )
    defaults.update(overrides)
    return argparse.Namespace(**defaults)


class TestMainFlow(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        (self.tmp / "conf.ini").write_text(CONFIG_TEMPLATE, encoding="utf-8")
        self.original_base_dir = main_module.get_base_dir
        self.original_build_client = main_module.build_client
        self.original_authorize = main_module.authorize
        main_module.get_base_dir = lambda: self.tmp

    def tearDown(self):
        main_module.get_base_dir = self.original_base_dir
        main_module.build_client = self.original_build_client
        main_module.authorize = self.original_authorize

    def _patch_client(self, client):
        main_module.build_client = lambda config, base_dir: client
        main_module.authorize = self._fake_authorize

    @staticmethod
    async def _fake_authorize(client, config, max_attempts=3):
        return None

    def _messages(self):
        paid = make_paid_reactions([types.MessageReactor(count=4,
                                                         peer_id=types.PeerUser(42))])
        return [
            make_message(msg_id=30, text="stars", reactions=paid),
            make_message(msg_id=29, text="no stars"),
            make_message(msg_id=28, text="older"),
        ]

    def test_full_run_creates_excel(self):
        client = FlowClient(messages=self._messages(),
                            entities={("user", 42): make_user(42, "ivan")})
        self._patch_client(client)

        code = asyncio.run(main_module.run(make_args()))

        self.assertEqual(code, main_module.EXIT_OK)
        files = list((self.tmp / "output").glob("*.xlsx"))
        self.assertEqual(len(files), 1)
        sheet = load_workbook(files[0])["Stars"]
        self.assertEqual(sheet.max_row, 2)            # шапка + одна запись
        column_of = {name: index + 1 for index, name in enumerate(COLUMNS)}
        self.assertEqual(sheet.cell(row=2, column=column_of["current_message_id"]).value, 30)
        self.assertEqual(sheet.cell(row=2, column=column_of["reactor_username"]).value, "ivan")
        self.assertEqual(sheet.cell(row=2, column=column_of["stars_count"]).value, 4)

    def test_run_without_stars_prints_message(self):
        client = FlowClient(messages=[make_message(msg_id=1, text="no stars")])
        self._patch_client(client)

        code = asyncio.run(main_module.run(make_args(count=1)))

        self.assertEqual(code, main_module.EXIT_OK)
        self.assertEqual(list((self.tmp / "output").glob("*.xlsx")), [])

    def test_interrupt_saves_partial_data(self):
        client = InterruptingClient(messages=self._messages(),
                                    entities={("user", 42): make_user(42, "ivan")})
        self._patch_client(client)

        code = asyncio.run(main_module.run(make_args(count=10)))

        self.assertEqual(code, main_module.EXIT_INTERRUPTED)
        files = list((self.tmp / "output").glob("*.xlsx"))
        self.assertEqual(len(files), 1)
        sheet = load_workbook(files[0])["Stars"]
        self.assertEqual(sheet.max_row, 2)

    def test_cancellation_saves_partial_data(self):
        """Ctrl+C в реальной жизни — это отмена задачи, данные всё равно сохраняются."""
        client = CancellingClient(messages=self._messages(),
                                  entities={("user", 42): make_user(42, "ivan")})
        self._patch_client(client)

        code = asyncio.run(main_module.run(make_args(count=10)))

        self.assertEqual(code, main_module.EXIT_INTERRUPTED)
        files = list((self.tmp / "output").glob("*.xlsx"))
        self.assertEqual(len(files), 1)
        sheet = load_workbook(files[0])["Stars"]
        self.assertEqual(sheet.max_row, 2)
        column_of = {name: index + 1 for index, name in enumerate(COLUMNS)}
        self.assertEqual(sheet.cell(row=2, column=column_of["reactor_username"]).value, "ivan")

    def test_output_dir_from_argument(self):
        client = FlowClient(messages=self._messages(),
                            entities={("user", 42): make_user(42, "ivan")})
        self._patch_client(client)
        custom = self.tmp / "custom dir"
        code = asyncio.run(main_module.run(make_args(output_dir=str(custom))))
        self.assertEqual(code, main_module.EXIT_OK)
        self.assertEqual(len(list(custom.glob("*.xlsx"))), 1)

    def test_prompt_count_validation(self):
        answers = iter(["0", "-5", "abc", " 12 "])
        original_input = builtins.input
        builtins.input = lambda *args, **kwargs: next(answers)
        try:
            self.assertEqual(main_module.prompt_count(), 12)
        finally:
            builtins.input = original_input

    def test_prompt_count_preset(self):
        self.assertEqual(main_module.prompt_count(7), 7)


class TestChannelQueue(TestMainFlow):
    """Несколько каналов через запятую: очередь, ошибки и Ctrl+C."""

    def _queue_client(self, interrupt_on=None, cancelling=False):
        paid = make_paid_reactions([types.MessageReactor(count=4,
                                                         peer_id=types.PeerUser(42))])
        channels = {
            "first": (make_channel(11, title="First Channel", username="first"),
                      [make_message(msg_id=30, text="stars", reactions=paid)]),
            "second": (make_channel(22, title="Second Channel", username="second"),
                       [make_message(msg_id=70, text="stars 2", reactions=paid)]),
        }
        client_class = QueueCancellingClient if cancelling else QueueClient
        return client_class(channels=channels,
                            entities={("user", 42): make_user(42, "ivan")},
                            interrupt_on=interrupt_on)

    def _files(self, directory=None):
        return sorted((directory or self.tmp / "output").glob("*.xlsx"))

    @staticmethod
    def _channel_of(path):
        sheet = load_workbook(path)["Stars"]
        column_of = {name: index + 1 for index, name in enumerate(COLUMNS)}
        return sheet.cell(row=2, column=column_of["current_channel"]).value

    def test_queue_parses_all_channels_in_order(self):
        client = self._queue_client()
        self._patch_client(client)

        code = asyncio.run(main_module.run(make_args(channel="@first, t.me/second")))

        self.assertEqual(code, main_module.EXIT_OK)
        self.assertEqual(client.parsed_order, ["first", "second"])
        files = self._files()
        self.assertEqual(len(files), 2)
        self.assertEqual([self._channel_of(path) for path in files],
                         ["First Channel", "Second Channel"])

    def test_queue_skips_unresolvable_channel(self):
        client = self._queue_client()
        self._patch_client(client)

        code = asyncio.run(main_module.run(make_args(channel="@first, @missing, @second")))

        self.assertEqual(code, main_module.EXIT_OK)
        self.assertEqual(client.parsed_order, ["first", "second"])
        self.assertEqual(len(self._files()), 2)

    def test_queue_interrupt_saves_current_channel_only(self):
        client = self._queue_client(interrupt_on="second")
        self._patch_client(client)

        code = asyncio.run(main_module.run(make_args(channel="@first, @second", count=10)))

        self.assertEqual(code, main_module.EXIT_INTERRUPTED)
        files = self._files()
        self.assertEqual(len(files), 1)
        self.assertEqual(self._channel_of(files[0]), "First Channel")

    def test_queue_cancellation_saves_current_channel_only(self):
        client = self._queue_client(interrupt_on="second", cancelling=True)
        self._patch_client(client)

        code = asyncio.run(main_module.run(make_args(channel="@first, @second", count=10)))

        self.assertEqual(code, main_module.EXIT_INTERRUPTED)
        files = self._files()
        self.assertEqual(len(files), 1)
        self.assertEqual(self._channel_of(files[0]), "First Channel")

    def test_prompt_channels_returns_queue_in_order(self):
        client = self._queue_client()
        tasks = asyncio.run(main_module.prompt_channels(client, "@second, @first, @second"))
        self.assertEqual([task.label for task in tasks], ["Second Channel", "First Channel"])

    def test_prompt_channels_reprompts_when_nothing_resolved(self):
        client = self._queue_client()
        original_input = builtins.input
        builtins.input = lambda *args, **kwargs: "@first"
        try:
            tasks = asyncio.run(main_module.prompt_channels(client, "@missing, @gone"))
        finally:
            builtins.input = original_input
        self.assertEqual([task.label for task in tasks], ["First Channel"])

    def test_prompt_channels_reprompts_on_bad_syntax(self):
        client = self._queue_client()
        answers = iter(["@first, @", "@first"])
        original_input = builtins.input
        builtins.input = lambda *args, **kwargs: next(answers)
        try:
            tasks = asyncio.run(main_module.prompt_channels(client, None))
        finally:
            builtins.input = original_input
        self.assertEqual([task.label for task in tasks], ["First Channel"])

    def test_queue_continues_after_channel_error(self):
        paid = make_paid_reactions([types.MessageReactor(count=4,
                                                         peer_id=types.PeerUser(42))])
        channels = {
            "first": (make_channel(11, title="First Channel", username="first"),
                      [make_message(msg_id=30, text="stars", reactions=paid)]),
            "second": (make_channel(22, title="Second Channel", username="second"),
                       [make_message(msg_id=70, text="stars 2", reactions=paid)]),
        }
        client = FailingChannelClient(channels=channels, fail_on={"second"},
                                      entities={("user", 42): make_user(42, "ivan")})
        self._patch_client(client)

        code = asyncio.run(main_module.run(make_args(channel="@second, @first")))

        self.assertEqual(code, main_module.EXIT_OK)
        self.assertEqual(client.parsed_order, ["second", "first"])
        files = self._files()
        self.assertEqual(len(files), 1)
        self.assertEqual(self._channel_of(files[0]), "First Channel")

    def test_queue_all_channels_failed(self):
        paid = make_paid_reactions([types.MessageReactor(count=4,
                                                         peer_id=types.PeerUser(42))])
        channels = {
            "first": (make_channel(11, title="First Channel", username="first"),
                      [make_message(msg_id=30, text="stars", reactions=paid)]),
            "second": (make_channel(22, title="Second Channel", username="second"),
                       [make_message(msg_id=70, text="stars 2", reactions=paid)]),
        }
        client = FailingChannelClient(channels=channels, fail_on={"first", "second"},
                                      entities={("user", 42): make_user(42, "ivan")})
        self._patch_client(client)

        code = asyncio.run(main_module.run(make_args(channel="@first, @second")))

        self.assertEqual(code, main_module.EXIT_ERROR)
        self.assertEqual(self._files(), [])

    def test_queue_exit_code(self):
        done = main_module.ChannelOutcome(task=main_module.ChannelTask(value="a", name="A"))
        failed = main_module.ChannelOutcome(task=main_module.ChannelTask(value="b", name="B"),
                                            status=main_module.STATUS_ERROR, error="boom")
        self.assertEqual(main_module.queue_exit_code([done], False), main_module.EXIT_OK)
        self.assertEqual(main_module.queue_exit_code([done, failed], False), main_module.EXIT_OK)
        self.assertEqual(main_module.queue_exit_code([failed], False), main_module.EXIT_ERROR)
        self.assertEqual(main_module.queue_exit_code([done], True),
                         main_module.EXIT_INTERRUPTED)

    def test_summary_counts_all_channels(self):
        client = self._queue_client()
        self._patch_client(client)
        asyncio.run(main_module.run(make_args(channel="@first, @second")))
        self.assertEqual(len(self._files()), 2)


if __name__ == "__main__":
    unittest.main(verbosity=2)
