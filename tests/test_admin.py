"""Тесты админки бота: добавление пользователей по id, whitelist, фикс статистики.

Проверяются:
* `UserStore` — хранение добавленных пользователей в users.json;
* `AccessMiddleware` — доступ для админов, статического whitelist и добавленных;
* хендлеры `/adduser`, `/deluser`, `/admin`, quick-add из уведомления;
* отсутствие в UI подсказки про `-100` и дублирующей строки `/help`;
* статистика «Обработано: N из M» больше не показывает технический лимит 1 000 000.
"""

from __future__ import annotations

import asyncio
import json
import sys
import tempfile
from pathlib import Path
from types import SimpleNamespace

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pytest

from bot.access import AccessMiddleware
from bot.config_bot import BotConfig
from bot.user_store import UserStore
from models import ParseStats


# --------------------------------------------------------------------------- #
# Заглушки aiogram-объектов
# --------------------------------------------------------------------------- #
class FakeBot:
    def __init__(self) -> None:
        self.sent: list[tuple[int, str]] = []
        self.documents: list = []

    async def send_message(self, chat_id, text, **kwargs):
        self.sent.append((chat_id, text))
        return SimpleNamespace(chat_id=chat_id, message_id=len(self.sent))

    async def send_document(self, chat_id, document, **kwargs):
        self.documents.append(document)
        return SimpleNamespace(chat_id=chat_id, message_id=len(self.sent) + 1)

    async def delete_message(self, chat_id, message_id, **kwargs):
        return True


class FakeMessage:
    def __init__(self, text: str = "", user_id: int = 1, username: str = "admin",
                 bot: FakeBot | None = None) -> None:
        self.text = text
        self.bot = bot or FakeBot()
        self.chat = SimpleNamespace(id=user_id)
        self.from_user = SimpleNamespace(
            id=user_id, username=username, full_name="Admin", is_bot=False,
        )
        self.answers: list[str] = []

    async def answer(self, text, **kwargs):
        self.answers.append(text)
        return self

    async def edit_text(self, text, **kwargs):
        self.answers.append(text)
        return self

    async def edit_reply_markup(self, **kwargs):
        return self

    @property
    def html_text(self) -> str:
        return self.answers[-1] if self.answers else ""


class FakeCallback:
    def __init__(self, data: str, user_id: int = 1, message: FakeMessage | None = None) -> None:
        self.data = data
        self.message = message or FakeMessage()
        self.bot = self.message.bot
        self.from_user = SimpleNamespace(id=user_id, username="admin", full_name="Admin")
        self.alerts: list[str] = []

    async def answer(self, text: str | None = None, show_alert: bool = False):
        if text:
            self.alerts.append(text)
        return True


class FakeState:
    def __init__(self) -> None:
        self.state = None
        self.data: dict = {}

    async def set_state(self, state):
        self.state = state

    async def get_state(self):
        return self.state

    async def clear(self):
        self.state = None
        self.data = {}

    async def update_data(self, **kwargs):
        self.data.update(kwargs)

    async def get_data(self):
        return dict(self.data)


def _aiogram_message(monkeypatch, user_id: int, username: str, bot: FakeBot):
    """Настоящий aiogram Message: middleware проверяет isinstance, ответы собираем."""
    from datetime import datetime, timezone

    from aiogram.types import Chat, Message, User

    answers: list[str] = []

    async def fake_answer(self, text, **kwargs):
        answers.append(text)
        return self

    monkeypatch.setattr(Message, "answer", fake_answer)
    message = Message.model_construct(
        message_id=1,
        date=datetime.now(timezone.utc),
        chat=Chat.model_construct(id=user_id, type="private"),
        from_user=User.model_construct(id=user_id, username=username,
                                       first_name="Stranger", is_bot=False),
        text="/parse",
        bot=bot,
    )
    return message, answers


def _aiogram_callback(monkeypatch, user_id: int, bot: FakeBot):
    """Настоящий aiogram CallbackQuery: ответы (алерты) собираем."""
    from datetime import datetime, timezone

    from aiogram.types import CallbackQuery, Chat, Message, User

    answers: list[str] = []

    async def fake_answer(self, text=None, show_alert=False, **kwargs):
        if text:
            answers.append(text)
        return True

    monkeypatch.setattr(CallbackQuery, "answer", fake_answer)
    message = Message.model_construct(
        message_id=1,
        date=datetime.now(timezone.utc),
        chat=Chat.model_construct(id=user_id, type="private"),
        from_user=User.model_construct(id=user_id, first_name="Stranger", is_bot=False),
    )
    callback = CallbackQuery.model_construct(
        id="1", chat_instance="1", data="wiz:status",
        from_user=User.model_construct(id=user_id, first_name="Stranger", is_bot=False),
        message=message,
        bot=bot,
    )
    return callback, answers


def _config(admins: set[int] | None = None, allowed: set[int] | None = None) -> BotConfig:
    return BotConfig(
        token="1:test",
        allowed_user_ids=allowed or set(),
        admin_ids=admins or {1},
    )


# --------------------------------------------------------------------------- #
# UserStore
# --------------------------------------------------------------------------- #
def test_user_store_add_remove_and_persist():
    with tempfile.TemporaryDirectory() as tmpdir:
        path = Path(tmpdir) / "users.json"
        store = UserStore(path)

        created, user = store.add(555, username="@vasya", added_by=1)
        assert created is True
        assert user.username == "vasya"
        assert 555 in store
        assert store.ids == {555}
        assert path.exists()

        # Повторное добавление не создаёт дубликат
        created_again, _ = store.add(555, full_name="Вася", added_by=1)
        assert created_again is False
        assert len(store) == 1
        assert store.get(555).full_name == "Вася"

        # Перезагрузка из файла
        reloaded = UserStore(path)
        assert reloaded.ids == {555}
        assert reloaded.get(555).username == "vasya"

        # Удаление
        assert reloaded.remove(555) is True
        assert reloaded.remove(555) is False
        assert UserStore(path).ids == set()


def test_user_store_ignores_corrupted_file():
    with tempfile.TemporaryDirectory() as tmpdir:
        path = Path(tmpdir) / "users.json"
        path.write_text("{не json", encoding="utf-8")
        store = UserStore(path)
        assert store.ids == set()

        path.write_text(json.dumps({"users": {"bad": {"username": "x"}, "42": {}}}), encoding="utf-8")
        store = UserStore(path)
        assert store.ids == {42}


# --------------------------------------------------------------------------- #
# AccessMiddleware
# --------------------------------------------------------------------------- #
def _run(coro):
    return asyncio.run(coro)


def test_access_middleware_roles_and_denial(monkeypatch):
    with tempfile.TemporaryDirectory() as tmpdir:
        store = UserStore(Path(tmpdir) / "users.json")
        store.add(333, added_by=1)
        middleware = AccessMiddleware({222}, admins={111}, store=store)

        assert middleware.is_allowed(111)  # админ
        assert middleware.is_allowed(222)  # из conf.ini
        assert middleware.is_allowed(333)  # добавлен через бота
        assert not middleware.is_allowed(444)

        calls: list[str] = []

        async def handler(event, data):
            calls.append("handled")

        # Разрешённый пользователь попадает в хендлер
        _run(middleware(handler, FakeMessage(user_id=333), {}))
        assert calls == ["handled"]

        # Незнакомый — не попадает, но админ получает уведомление с кнопкой
        admin_bot = FakeBot()
        event, answers = _aiogram_message(monkeypatch, user_id=444, username="stranger",
                                          bot=admin_bot)
        _run(middleware(handler, event, {"bot": admin_bot}))
        assert calls == ["handled"]  # хендлер не вызван
        assert any("Доступ запрещён" in text for text in answers)
        assert admin_bot.sent and admin_bot.sent[0][0] == 111
        assert "444" in admin_bot.sent[0][1]


def test_access_middleware_callback_denied(monkeypatch):
    middleware = AccessMiddleware(set(), admins={111})
    bot = FakeBot()
    callback, answers = _aiogram_callback(monkeypatch, user_id=444, bot=bot)

    async def handler(event, data):
        raise AssertionError("не должен вызываться")

    _run(middleware(handler, callback, {"bot": bot}))
    # Хендлер не вызван, пользователю отвечено алертом, админ уведомлён
    assert answers and "Доступ запрещён" in answers[0]
    assert bot.sent and bot.sent[0][0] == 111
    assert "444" in bot.sent[0][1]


# --------------------------------------------------------------------------- #
# Админ-хендлеры
# --------------------------------------------------------------------------- #
def test_admin_add_and_remove_user_flow():
    import bot.handlers_bot as hb

    with tempfile.TemporaryDirectory() as tmpdir:
        store = UserStore(Path(tmpdir) / "users.json")
        config = _config(admins={1})
        state = FakeState()

        # /adduser без аргументов — админ вводит id отдельным сообщением
        msg = FakeMessage(text="/adduser")
        _run(hb.cmd_adduser(msg, state, config, store))
        assert state.state == hb.AdminState.ADD_USER
        assert "id" in msg.answers[-1].lower()

        # Пришел id (можно несколько через запятую, с пометкой @username)
        msg2 = FakeMessage(text="777, 888 @vasya")
        state.data = {}
        _run(hb.process_admin_add(msg2, state, config, store))
        assert store.ids == {777, 888}
        assert store.get(888).username == "vasya"
        assert state.state is None  # состояние сброшено
        # добавленному пользователю отправлено уведомление о доступе
        assert any(chat_id == 777 for chat_id, _ in msg2.bot.sent)

        # /adduser с id сразу в команде
        msg3 = FakeMessage(text="/adduser 999")
        _run(hb.cmd_adduser(msg3, state, config, store))
        assert 999 in store

        # /deluser
        msg4 = FakeMessage(text="/deluser 777")
        _run(hb.cmd_deluser(msg4, state, config, store))
        assert 777 not in store

        # Админа удалить нельзя, а id из conf.ini — только через conf.ini
        config2 = _config(admins={1}, allowed={222})
        msg5 = FakeMessage(text="/deluser 1 222 999")
        _run(hb.cmd_deluser(msg5, state, config2, store))
        assert "администратор" in msg5.answers[-1]
        assert "conf.ini" in msg5.answers[-1]
        assert 999 not in store


def test_admin_panel_requires_admin_rights():
    import bot.handlers_bot as hb

    with tempfile.TemporaryDirectory() as tmpdir:
        store = UserStore(Path(tmpdir) / "users.json")
        config = _config(admins={1})

        # Обычный пользователь не может открыть панель
        msg = FakeMessage(text="/admin", user_id=999)
        _run(hb.cmd_admin(msg, FakeState(), config, store))
        assert "только администратору" in msg.answers[-1]

        # Админ — видит панель с кнопкой добавления
        admin_msg = FakeMessage(text="/admin", user_id=1)
        _run(hb.cmd_admin(admin_msg, FakeState(), config, store))
        assert "Доступ к боту" in admin_msg.answers[-1]

        callback = FakeCallback("adm:add", user_id=1)
        _run(hb.cb_admin_add(callback, FakeState(), config))
        assert callback.message.answers  # инструкция отправлена


def test_admin_quick_add_from_notification():
    import bot.handlers_bot as hb

    with tempfile.TemporaryDirectory() as tmpdir:
        store = UserStore(Path(tmpdir) / "users.json")
        config = _config(admins={1})
        callback = FakeCallback("adm:add:555", user_id=1)

        _run(hb.cb_admin_quick_add(callback, config, store))
        assert store.ids == {555}
        assert callback.alerts == ["Доступ открыт"]

        # Повторное нажатие не дублирует запись
        callback2 = FakeCallback("adm:add:555", user_id=1)
        _run(hb.cb_admin_quick_add(callback2, config, store))
        assert store.ids == {555}
        assert callback2.alerts == ["Уже в списке"]


# --------------------------------------------------------------------------- #
# UI-тексты
# --------------------------------------------------------------------------- #
def test_ui_texts_have_no_hint_and_no_help_line():
    import bot.handlers_bot as hb

    help_text = hb._help_text(admin=True)
    assert "/help" not in help_text  # строка «/help — эта справка» убрана
    assert "Бот для пробива донатеров звёзд в Telegram" in hb.BOT_TITLE

    channels_step = hb._channels_step_text()
    assert "без -100" not in channels_step
    assert "-100" not in channels_step

    # /start больше не подписан как «Бот для сбора звёзд (Paid Reactions)»
    assert "Paid Reactions" not in hb.BOT_TITLE


def test_start_message_for_admin_has_admin_button():
    import bot.handlers_bot as hb

    msg_admin = FakeMessage(text="/start", user_id=1)
    _run(hb.cmd_start(msg_admin, FakeState(), _config(admins={1})))
    assert "пробива донатеров" in msg_admin.answers[-1]

    msg_user = FakeMessage(text="/start", user_id=1)
    _run(hb.cmd_start(msg_user, FakeState(), _config(admins=set())))
    assert "пробива донатеров" in msg_user.answers[-1]


# --------------------------------------------------------------------------- #
# Статистика: никаких «из 1000000»
# --------------------------------------------------------------------------- #
def test_stats_unlimited_and_exhausted():
    stats = ParseStats()
    stats.set_requested(1_000_000, unlimited=True)
    assert stats.requested is None  # технический лимит в отчёт не попадает

    stats.scanned = 6957
    stats.exhausted = True
    stats.finalize_requested()
    assert stats.requested == 6957

    limited = ParseStats()
    limited.set_requested(500)
    limited.scanned = 500
    limited.exhausted = True   # дошли до конца ровно на лимите
    limited.finalize_requested()
    assert limited.requested == 500

    interrupted = ParseStats()
    interrupted.set_requested(500)
    interrupted.scanned = 120
    interrupted.interrupted = True
    interrupted.finalize_requested()
    assert interrupted.requested == 500  # прерывание не меняет план


def test_iter_message_batches_marks_exhausted():
    """Конец истории канала помечается как exhausted, чтобы план стал точным."""
    from datetime import datetime, timezone

    from handlers import iter_message_batches

    class DummyClient:
        def __init__(self, msgs):
            self.msgs = msgs

        async def iter_messages(self, entity, limit=None, offset_id=0, **kwargs):
            for m in self.msgs:
                if offset_id and m.id >= offset_id:
                    continue
                yield m

    def msg(mid: int):
        return SimpleNamespace(id=mid, date=datetime(2024, 1, mid, tzinfo=timezone.utc), reactions=None)

    client = DummyClient([msg(3), msg(2), msg(1)])
    stats = ParseStats()
    stats.set_requested(1_000_000, unlimited=True)

    async def collect():
        gathered = []
        async for batch in iter_message_batches(
            client=client, entity="dummy", limit=1_000_000, batch_size=10, stats=stats,
        ):
            gathered.extend(batch)
        return gathered

    collected = _run(collect())
    assert len(collected) == 3
    assert stats.exhausted is True

    stats.scanned = len(collected)
    stats.finalize_requested()
    assert stats.requested == 3


def test_cli_excel_username_rule():
    """Лист Stars CLI: username при наличии, «отсутствует» при отсутствии."""
    from exporter import build_dataframe
    from models import REACTOR_ANONYMOUS, REACTOR_USER, USERNAME_MISSING, StarRecord

    records = [
        StarRecord("text", REACTOR_USER, "@ch", 1, "@ch", 1, "ivan", 7, 10,
                   reactor_has_username=True),
        StarRecord("text", REACTOR_USER, "@ch", 2, "@ch", 2, "Ivan Petrov", 8, 5),
        StarRecord("text", REACTOR_ANONYMOUS, "@ch", 3, "@ch", 3, "", None, 4),
    ]
    frame = build_dataframe(records)
    usernames = list(frame["reactor_username"])
    assert usernames == ["ivan", USERNAME_MISSING, USERNAME_MISSING]
    assert frame.loc[0, "post_link"] == ""


def test_bot_build_donation_rows_fields():
    """Лист «Донаты»: состав и порядок полей 1:1 как на согласованном образце."""
    from exporter import DONATION_COLUMNS, build_donation_rows
    from models import COLUMNS

    assert DONATION_COLUMNS == COLUMNS
    assert list(DONATION_COLUMNS) == [
        "message_type", "reactor_type", "current_channel", "current_message_id",
        "post_link", "original_channel", "original_message_id", "reactor_username",
        "reactor_id", "stars_count",
    ]

    from models import REACTOR_USER, StarRecord

    rows = build_donation_rows([
        StarRecord("photo", REACTOR_USER, "Канал", 100, "Канал", 100, "vasya", 5, 3,
                   "https://t.me/kanal/100", reactor_has_username=True),
    ])
    assert rows[0]["message_type"] == "photo"
    assert rows[0]["post_link"] == "https://t.me/kanal/100"
    assert rows[0]["reactor_username"] == "vasya"
    assert rows[0]["stars_count"] == 3


# --------------------------------------------------------------------------- #
# Сквозной поток: wizard -> отчёт в Excel
# --------------------------------------------------------------------------- #
def test_wizard_start_button_and_scope_modes():
    import bot.handlers_bot as hb
    from bot.queue_service import QueueService

    config = _config(admins={1})
    queue = QueueService()  # воркер не запускаем: задачи только ставятся в очередь
    state = FakeState()

    # Кнопка «🚀 Запустить парсинг»
    callback = FakeCallback("wiz:start", user_id=1)
    _run(hb.cb_wizard_start(callback, state, queue, config))
    assert state.state == hb.WizardState.CHANNELS

    # Шаг «Все посты» -> режим «все донатеры»: лимит не должен быть планом отчёта
    callback_all = FakeCallback("scope_all", user_id=1)
    _run(hb.cb_scope_all(callback_all, state, config))
    assert state.state == hb.WizardState.MODE
    assert state.data["unlimited"] is True
    assert state.data["limit"] == hb.UNLIMITED_LIMIT
    assert state.data["scope_desc"] == "Все посты"


def test_final_results_excel_and_unlimited_stats(tmp_path, monkeypatch):
    """_send_final_results: лист «Донаты» с полями образца, статистика без «из 1000000»."""
    import openpyxl

    import bot.handlers_bot as hb
    from exporter import DONATION_COLUMNS
    from models import REACTOR_USER, StarRecord

    monkeypatch.chdir(tmp_path)
    stats = {"Канал": ParseStats(requested=None, scanned=6957, exhausted=True,
                                 with_stars=1, reactors=1, not_found=1)}
    records = [
        StarRecord("photo", REACTOR_USER, "Канал", 100, "Канал", 100, "vasya", 5, 3,
                   "https://t.me/kanal/100", reactor_has_username=True),
    ]
    bot = FakeBot()
    _run(hb._send_final_results(
        bot=bot,
        chat_id=1,
        status_msg_id=42,
        all_records=records,
        stats_by_channel=stats,
        channel_outcomes={"Канал": "успешно"},
        scope_desc="Все посты",
        only_above_threshold=False,
        bot_config=_config(admins={1}),
        interrupted=False,
        channels_meta={"Канал": {}},
        unlimited=True,
    ))

    files = list((tmp_path / "output").glob("*.xlsx"))
    assert files, "файл отчёта должен быть создан"
    wb = openpyxl.load_workbook(files[0])
    assert wb.sheetnames == ["Донаты", "Сводка", "Инфо"]
    assert [cell.value for cell in wb["Донаты"][1]] == list(DONATION_COLUMNS)
    assert wb["Донаты"].cell(row=2, column=DONATION_COLUMNS.index("reactor_username") + 1).value == "vasya"

    # В сообщениях чата нет технического лимита
    chat_text = "\n".join(text for _chat_id, text in bot.sent)
    assert "1000000" not in chat_text and "1 000 000" not in chat_text
    assert "Обработано:</b> 6957" in chat_text
    assert bot.documents, "xlsx должен быть отправлен в чат"


def test_status_message_without_limit():
    import bot.handlers_bot as hb
    from bot.queue_service import ParseTask, QueueService, STATUS_RUNNING, TASK_TYPE_PARSE

    queue = QueueService()
    task = ParseTask(
        task_id="t1", user_id=1, user_display="@admin", task_type=TASK_TYPE_PARSE,
        channels=["@ch"], scope_desc="Все посты",
    )
    task.status = STATUS_RUNNING
    task.current_channel = "@ch"
    task.scanned = 6957
    task.total_hint = None
    queue.tasks.append(task)

    message = FakeMessage(user_id=1)
    _run(hb._answer_status(message, 1, queue))
    assert "6957" in message.answers[-1]
    assert "1000000" not in message.answers[-1]
    assert "/1000000" not in message.answers[-1]
