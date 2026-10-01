"""Офлайн-тесты: санитайзер имён, конфиг, определение типов, реакторов, экспорт.

Запуск (без Telegram и без сети):
    python -m unittest discover -s tests -v
    # или
    python tests/test_offline.py
"""

from __future__ import annotations

import asyncio
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook  # noqa: E402
from telethon.errors import FloodWaitError  # noqa: E402
from telethon.tl import types  # noqa: E402

from client import ConnectionFailureError, connect  # noqa: E402
from config import ConfigError, load_config  # noqa: E402
from exporter import build_result_filename, export_records, safe_filename, unique_path, write_excel  # noqa: E402
from handlers import (  # noqa: E402
    ChannelResolutionError,
    EntityCache,
    build_records_for_message,
    channel_input_text,
    channel_title,
    detect_message_type,
    get_paid_reactions_total_count,
    has_paid_reactions,
    message_type_with_forward,
    parse_channel,
    parse_channel_input,
    parse_channels_input,
    resolve_reactor,
    split_channel_list,
)
from models import COLUMNS, NOT_FOUND, REACTOR_ANONYMOUS, REACTOR_CHANNEL, REACTOR_USER, StarRecord  # noqa: E402


# --------------------------------------------------------------------------- #
# Фейковый клиент
# --------------------------------------------------------------------------- #
class FakeClient:
    """Мини-замена TelegramClient: только get_entity и iter_messages."""

    def __init__(self, messages=None, entities=None, missing=(), flood_once=False):
        self.messages = list(messages or [])
        self.entities = dict(entities or {})
        self.missing = set(missing)
        self.flood_once = flood_once
        self.entity_calls = 0
        self.iter_calls = 0

    async def get_entity(self, peer):
        self.entity_calls += 1
        key = ("user", peer.user_id) if isinstance(peer, types.PeerUser) else (
            "channel", peer.channel_id) if isinstance(peer, types.PeerChannel) else ("other", 0)
        if key in self.missing:
            raise ValueError(f"Could not find the input entity for {peer}")
        return self.entities.get(key)

    async def iter_messages(self, entity, limit=None, offset_id=0, **kwargs):
        self.iter_calls += 1
        offset_id = offset_id or 0
        selected = [m for m in self.messages if not offset_id or int(m.id) < offset_id]
        if limit:
            selected = selected[:limit]
        for index, message in enumerate(selected):
            yield message
            if self.flood_once and index == 0:
                self.flood_once = False
                raise FloodWaitError(request=None, capture=0)


def make_message(msg_id=1, text=None, media=None, poll=None, fwd_from=None, reactions=None):
    """Простой объект сообщения (достаточно для используемых атрибутов)."""
    return SimpleNamespace(
        id=msg_id, message=text, media=media, poll=poll,
        fwd_from=fwd_from, reactions=reactions,
    )


def make_paid_reactions(reactors):
    """Объект reactions с платными реакциями и top_reactors."""
    return types.MessageReactions(
        results=[types.ReactionCount(reaction=types.ReactionPaid(), count=sum(
            getattr(r, "count", 0) for r in reactors))],
        top_reactors=reactors,
    )


def make_user(user_id, username=None, first="Ivan", last="Petrov"):
    return types.User(id=user_id, username=username, first_name=first, last_name=last,
                      access_hash=123)


def make_channel(channel_id, title="Test Channel", username="testchan"):
    return types.Channel(id=channel_id, title=title, photo=types.ChatPhotoEmpty(),
                         date=datetime(2024, 1, 1), access_hash=1, username=username,
                         broadcast=True)


# --------------------------------------------------------------------------- #
# Имя файла
# --------------------------------------------------------------------------- #
class TestSafeFilename(unittest.TestCase):
    def test_illegal_characters_replaced(self):
        self.assertEqual(safe_filename('SupernovaElit Premium|Chat️'),
                         "SupernovaElit Premium_Chat")

    def test_all_illegal_characters(self):
        self.assertEqual(safe_filename('a\\b/c:d*e?f"g<h>i|j'), "a_b_c_d_e_f_g_h_i_j")

    def test_invisible_and_variation_selectors_removed(self):
        self.assertEqual(safe_filename("Cha️nnel​‍﻿"), "Channel")

    def test_trailing_spaces_and_dots(self):
        self.assertEqual(safe_filename("  my channel.  "), "my channel")
        self.assertEqual(safe_filename("channel..."), "channel")

    def test_empty_falls_back(self):
        self.assertEqual(safe_filename(""), "channel")
        self.assertEqual(safe_filename("️‍"), "channel")
        self.assertEqual(safe_filename("///"), "channel")

    def test_reserved_windows_names(self):
        self.assertEqual(safe_filename("CON"), "_CON")
        self.assertEqual(safe_filename("nul.txt"), "_nul.txt")

    def test_max_length(self):
        name = safe_filename("A" * 500, max_length=30)
        self.assertEqual(len(name), 30)
        self.assertEqual(name, "A" * 30)
        # Обрезка не должна оставлять точку/пробел в конце
        self.assertEqual(safe_filename("B" * 40 + " .", max_length=20), "B" * 20)

    def test_result_filename_format(self):
        name = build_result_filename("My Channel", datetime(2025, 3, 4, 5, 6, 7))
        self.assertEqual(name, "2025-03-04_05-06-07_My Channel.xlsx")
        self.assertNotIn("/", name)
        self.assertNotIn(":", name)

    def test_result_filename_long_channel(self):
        name = build_result_filename("C" * 300, datetime(2025, 1, 1), max_length=50)
        self.assertTrue(name.startswith("2025-01-01_00-00-00_"))
        self.assertLessEqual(len(name), len("2025-01-01_00-00-00_") + 50 + len(".xlsx"))

    def test_unique_path(self):
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp)
            (directory / "a.xlsx").write_text("x")
            first = unique_path(directory, "a.xlsx")
            self.assertEqual(first.name, "a_1.xlsx")
            (directory / "a_1.xlsx").write_text("x")
            self.assertEqual(unique_path(directory, "a.xlsx").name, "a_2.xlsx")


# --------------------------------------------------------------------------- #
# Конфигурация
# --------------------------------------------------------------------------- #
SAMPLE_CONFIG = """
[API]
ID_API=123456
HASH_API=0123456789abcdef0123456789abcdef
PHONE=+79991234567
CLOUD_PASSWORD=secret

[SESSION]
NAME=my_session
DEVICE_MODEL=SM-G991B
SYSTEM_VERSION=Android 14
APP_VERSION=8.8.2
LANG_CODE=en

[DELAY]
MESSAGES_INTERVAL_MIN=0.05
MESSAGES_INTERVAL_MAX=0.1
"""


class TestConfig(unittest.TestCase):
    def _write(self, content: str) -> Path:
        directory = Path(tempfile.mkdtemp())
        path = directory / "conf.ini"
        path.write_text(content, encoding="utf-8")
        return path

    def test_valid_config(self):
        path = self._write(SAMPLE_CONFIG)
        config = load_config(path.parent, path)
        self.assertEqual(config.api.api_id, 123456)
        self.assertEqual(config.api.cloud_password, "secret")
        self.assertEqual(config.session.name, "my_session")
        self.assertEqual(config.delay.as_tuple(), (0.05, 0.1))
        self.assertEqual(config.output.directory, "output")
        self.assertFalse(config.parser.warmup_participants)

    def test_missing_file(self):
        with self.assertRaises(ConfigError):
            load_config(Path(tempfile.mkdtemp()))

    def test_missing_api_id(self):
        path = self._write(SAMPLE_CONFIG.replace("ID_API=123456", "ID_API="))
        with self.assertRaises(ConfigError):
            load_config(path.parent, path)

    def test_non_integer_api_id(self):
        path = self._write(SAMPLE_CONFIG.replace("ID_API=123456", "ID_API=12a34"))
        with self.assertRaises(ConfigError):
            load_config(path.parent, path)

    def test_empty_phone_and_hash(self):
        path = self._write(SAMPLE_CONFIG.replace("PHONE=+79991234567", "PHONE="))
        with self.assertRaises(ConfigError):
            load_config(path.parent, path)
        path = self._write(SAMPLE_CONFIG.replace("HASH_API=0123456789abcdef0123456789abcdef",
                                                  "HASH_API="))
        with self.assertRaises(ConfigError):
            load_config(path.parent, path)

    def test_empty_cloud_password_is_none(self):
        path = self._write(SAMPLE_CONFIG.replace("CLOUD_PASSWORD=secret", "CLOUD_PASSWORD="))
        config = load_config(path.parent, path)
        self.assertIsNone(config.api.cloud_password)

    def test_delay_bounds_swapped(self):
        path = self._write(SAMPLE_CONFIG
                           .replace("MESSAGES_INTERVAL_MIN=0.05", "MESSAGES_INTERVAL_MIN=1.5")
                           .replace("MESSAGES_INTERVAL_MAX=0.1", "MESSAGES_INTERVAL_MAX=0.5"))
        config = load_config(path.parent, path)
        self.assertEqual(config.delay.as_tuple(), (0.5, 1.5))

    def test_case_insensitive_sections_and_keys(self):
        path = self._write("""
[api]
id_api=999
hash_api=abcabc
phone=+79991112233
cloud_password=

[delay]
messages_interval_min=0.2
messages_interval_max=1
""")
        config = load_config(path.parent, path)
        self.assertEqual(config.api.api_id, 999)
        self.assertEqual(config.delay.as_tuple(), (0.2, 1.0))
        self.assertEqual(config.session.name, "anon")

    def test_invalid_delay(self):
        path = self._write(SAMPLE_CONFIG.replace("MESSAGES_INTERVAL_MAX=0.1",
                                                 "MESSAGES_INTERVAL_MAX=abc"))
        with self.assertRaises(ConfigError):
            load_config(path.parent, path)


# --------------------------------------------------------------------------- #
# Ввод канала
# --------------------------------------------------------------------------- #
class TestChannelInput(unittest.TestCase):
    def test_variants(self):
        cases = {
            "durov": "durov",
            "@durov": "durov",
            "https://t.me/durov": "durov",
            "t.me/durov": "durov",
            "http://www.telegram.me/durov/123": "durov",
            "t.me/s/durov": "durov",
            "-1001234567890": -1001234567890,
            "1234567890": 1234567890,
        }
        for raw, expected in cases.items():
            self.assertEqual(parse_channel_input(raw), expected, raw)

    def test_c_link(self):
        self.assertEqual(parse_channel_input("https://t.me/c/1234567890"), -1001234567890)
        self.assertEqual(parse_channel_input("https://t.me/c/1001234567890"), -1001234567890)

    def test_invite_links(self):
        self.assertEqual(parse_channel_input("https://t.me/+abcdefghijklmnopq"),
                         "https://t.me/+abcdefghijklmnopq")
        self.assertEqual(parse_channel_input("t.me/joinchat/AAAAAEHbEkejzxUjAUCfYg"),
                         "https://t.me/joinchat/AAAAAEHbEkejzxUjAUCfYg")

    def test_empty(self):
        with self.assertRaises(Exception):
            parse_channel_input("   ")


# --------------------------------------------------------------------------- #
# Несколько каналов через запятую (очередь)
# --------------------------------------------------------------------------- #
class TestChannelsInput(unittest.TestCase):
    def test_comma_separated(self):
        self.assertEqual(parse_channels_input("@durov, t.me/telegram, -1001234567890"),
                         ["durov", "telegram", -1001234567890])

    def test_order_is_preserved(self):
        self.assertEqual(parse_channels_input("@c, @a, @b"), ["c", "a", "b"])

    def test_semicolon_and_newline_are_separators_too(self):
        self.assertEqual(parse_channels_input("@a;@b\n@c"), ["a", "b", "c"])

    def test_duplicates_removed_case_insensitive(self):
        self.assertEqual(parse_channels_input("@durov, durov, t.me/DUROV"), ["durov"])

    def test_empty_parts_are_ignored(self):
        self.assertEqual(parse_channels_input("  @a ,, , @b ,  "), ["a", "b"])

    def test_invite_links_survive_split(self):
        self.assertEqual(parse_channels_input("t.me/+abcdefghijklmnopq, @durov"),
                         ["https://t.me/+abcdefghijklmnopq", "durov"])

    def test_only_separators(self):
        with self.assertRaises(ChannelResolutionError):
            parse_channels_input("  ,, ; ")

    def test_bad_item_reports_which_one(self):
        with self.assertRaises(ChannelResolutionError) as ctx:
            parse_channels_input("@durov, @")
        self.assertIn("'@'", str(ctx.exception))

    def test_split_channel_list(self):
        self.assertEqual(split_channel_list(" @a , b ;\nc "), ["@a", "b", "c"])
        self.assertEqual(split_channel_list(""), [])

    def test_channel_input_text(self):
        self.assertEqual(channel_input_text("durov"), "@durov")
        self.assertEqual(channel_input_text(-1001234567890), "-1001234567890")
        self.assertEqual(channel_input_text("https://t.me/+abcdefghijklmnopq"),
                         "https://t.me/+abcdefghijklmnopq")


# --------------------------------------------------------------------------- #
# Типы сообщений
# --------------------------------------------------------------------------- #
class TestMessageType(unittest.TestCase):
    def test_text(self):
        self.assertEqual(detect_message_type(make_message(text="hello")), "text")
        self.assertEqual(detect_message_type(make_message(text="   ")), "other")

    def test_photo(self):
        message = make_message(media=types.MessageMediaPhoto(photo=types.PhotoEmpty(id=1)))
        self.assertEqual(detect_message_type(message), "photo")

    def test_webpage(self):
        message = make_message(media=types.MessageMediaWebPage(
            webpage=types.WebPageEmpty(id=1)))
        self.assertEqual(detect_message_type(message), "webpage")

    def test_poll(self):
        self.assertEqual(detect_message_type(make_message(poll=object())), "poll")

    def test_video(self):
        document = types.Document(id=1, access_hash=1, dc_id=1, file_reference=b"",
                                  date=datetime(2024, 1, 1), mime_type="video/mp4", size=10,
                                  attributes=[types.DocumentAttributeVideo(10, 100, 100,
                                                                           False, False)])
        message = make_message(media=types.MessageMediaDocument(document=document))
        self.assertEqual(detect_message_type(message), "video")

    def test_voice(self):
        document = types.Document(id=1, access_hash=1, dc_id=1, file_reference=b"",
                                  date=datetime(2024, 1, 1), mime_type="audio/ogg", size=10,
                                  attributes=[types.DocumentAttributeAudio(10, voice=True)])
        message = make_message(media=types.MessageMediaDocument(document=document))
        self.assertEqual(detect_message_type(message), "voice")

    def test_audio(self):
        document = types.Document(id=1, access_hash=1, dc_id=1, file_reference=b"",
                                  date=datetime(2024, 1, 1), mime_type="audio/mpeg", size=10,
                                  attributes=[types.DocumentAttributeAudio(10)])
        message = make_message(media=types.MessageMediaDocument(document=document))
        self.assertEqual(detect_message_type(message), "audio")

    def test_gif_priority(self):
        document = types.Document(id=1, access_hash=1, dc_id=1, file_reference=b"",
                                  date=datetime(2024, 1, 1), mime_type="video/mp4", size=10,
                                  attributes=[types.DocumentAttributeVideo(1, 10, 10, False, False),
                                              types.DocumentAttributeAnimated()])
        message = make_message(media=types.MessageMediaDocument(document=document))
        self.assertEqual(detect_message_type(message), "gif")

    def test_sticker(self):
        document = types.Document(id=1, access_hash=1, dc_id=1, file_reference=b"",
                                  date=datetime(2024, 1, 1), mime_type="image/webp", size=10,
                                  attributes=[types.DocumentAttributeSticker(alt=":)",
                                                                             stickerset=None)])
        message = make_message(media=types.MessageMediaDocument(document=document))
        self.assertEqual(detect_message_type(message), "sticker")

    def test_plain_document(self):
        document = types.Document(id=1, access_hash=1, dc_id=1, file_reference=b"",
                                  date=datetime(2024, 1, 1), mime_type="application/pdf",
                                  size=10, attributes=[types.DocumentAttributeFilename("a.pdf")])
        message = make_message(media=types.MessageMediaDocument(document=document))
        self.assertEqual(detect_message_type(message), "document")

    def test_other_media(self):
        message = make_message(media=types.MessageMediaGeo(
            geo=types.GeoPointEmpty()))
        self.assertEqual(detect_message_type(message), "other")

    def test_forward_suffix(self):
        self.assertEqual(message_type_with_forward("video", True), "video(forward)")
        self.assertEqual(message_type_with_forward("video", False), "video")


# --------------------------------------------------------------------------- #
# Реакторы и звёзды
# --------------------------------------------------------------------------- #
class TestReactors(unittest.TestCase):
    def setUp(self):
        self.cache = EntityCache()

    def run_async(self, coro):
        return asyncio.run(coro)

    def test_user_found(self):
        client = FakeClient(entities={("user", 42): make_user(42, username="ivan")})
        reactor = types.MessageReactor(count=10, peer_id=types.PeerUser(42))
        result = self.run_async(resolve_reactor(client, reactor, self.cache))
        self.assertEqual(result, (REACTOR_USER, 42, "ivan"))

    def test_user_without_username_uses_name(self):
        client = FakeClient(entities={("user", 7): make_user(7, username=None,
                                                             first="Ivan", last="Petrov")})
        reactor = types.MessageReactor(count=3, peer_id=types.PeerUser(7))
        self.assertEqual(self.run_async(resolve_reactor(client, reactor, self.cache))[2],
                         "Ivan Petrov")

    def test_user_not_found(self):
        client = FakeClient(missing={("user", 99)})
        reactor = types.MessageReactor(count=5, peer_id=types.PeerUser(99))
        result = self.run_async(resolve_reactor(client, reactor, self.cache))
        self.assertEqual(result, (REACTOR_USER, 99, NOT_FOUND))

    def test_channel_reactor(self):
        client = FakeClient(entities={("channel", 555): make_channel(555, title="News",
                                                                     username="news")})
        reactor = types.MessageReactor(count=50, peer_id=types.PeerChannel(555))
        result = self.run_async(resolve_reactor(client, reactor, self.cache))
        self.assertEqual(result, (REACTOR_CHANNEL, 555, "news"))

    def test_channel_without_username_uses_title(self):
        client = FakeClient(entities={("channel", 556): make_channel(556, title="News",
                                                                     username=None)})
        reactor = types.MessageReactor(count=1, peer_id=types.PeerChannel(556))
        self.assertEqual(self.run_async(resolve_reactor(client, reactor, self.cache))[2], "News")

    def test_anonymous(self):
        reactor = types.MessageReactor(count=25, anonymous=True)
        self.assertEqual(self.run_async(resolve_reactor(FakeClient(), reactor, self.cache)),
                         (REACTOR_ANONYMOUS, None, ""))

    def test_peer_none_is_anonymous(self):
        reactor = types.MessageReactor(count=25, peer_id=None)
        self.assertEqual(self.run_async(resolve_reactor(FakeClient(), reactor, self.cache)),
                         (REACTOR_ANONYMOUS, None, ""))

    def test_cache_avoids_second_request(self):
        client = FakeClient(entities={("user", 42): make_user(42, "ivan")})
        reactor = types.MessageReactor(count=1, peer_id=types.PeerUser(42))
        self.run_async(resolve_reactor(client, reactor, self.cache))
        self.run_async(resolve_reactor(client, reactor, self.cache))
        self.assertEqual(client.entity_calls, 1)

    def test_paid_reactions_detection(self):
        self.assertTrue(has_paid_reactions(make_paid_reactions([
            types.MessageReactor(count=1, peer_id=types.PeerUser(1))])))
        self.assertFalse(has_paid_reactions(types.MessageReactions(
            results=[types.ReactionCount(reaction=types.ReactionEmoji(emoticon="👍"), count=1)])))
        self.assertFalse(has_paid_reactions(None))
        self.assertFalse(has_paid_reactions(make_message().reactions))

    def test_paid_reactions_total_count(self):
        self.assertEqual(get_paid_reactions_total_count(None), 0)
        self.assertEqual(get_paid_reactions_total_count(types.MessageReactions(results=[])), 0)
        self.assertEqual(get_paid_reactions_total_count(types.MessageReactions(
            results=[types.ReactionCount(reaction=types.ReactionEmoji(emoticon="❤️"), count=5)]
        )), 0)
        self.assertEqual(get_paid_reactions_total_count(types.MessageReactions(
            results=[
                types.ReactionCount(reaction=types.ReactionEmoji(emoticon="❤️"), count=5),
                types.ReactionCount(reaction=types.ReactionPaid(), count=25),
            ]
        )), 25)

    def test_build_records_paid_without_reactors_fallback_anonymous(self):
        client = FakeClient()
        reactions = types.MessageReactions(
            results=[types.ReactionCount(reaction=types.ReactionPaid(), count=15)],
            top_reactors=[],
        )
        message = make_message(msg_id=101, text="exclusive post", reactions=reactions)
        records = self.run_async(
            build_records_for_message(client, message, "My Channel", self.cache)
        )
        self.assertEqual(len(records), 1)
        record = records[0]
        self.assertEqual(record.reactor_type, REACTOR_ANONYMOUS)
        self.assertEqual(record.reactor_id, None)
        self.assertEqual(record.reactor_username, "")
        self.assertEqual(record.stars_count, 15)
        self.assertEqual(record.current_message_id, 101)
        self.assertEqual(record.current_channel, "My Channel")

    def test_build_records_simple_post(self):
        client = FakeClient(entities={("user", 42): make_user(42, "ivan")})
        message = make_message(
            msg_id=100, text="hello",
            reactions=make_paid_reactions([types.MessageReactor(count=7,
                                                                peer_id=types.PeerUser(42))]),
        )
        records = self.run_async(
            build_records_for_message(client, message, "My Channel", self.cache))
        self.assertEqual(len(records), 1)
        record = records[0]
        self.assertEqual(record.message_type, "text")
        self.assertEqual(record.current_channel, "My Channel")
        self.assertEqual(record.current_message_id, 100)
        self.assertEqual(record.original_channel, "My Channel")   # дубль текущего канала
        self.assertEqual(record.original_message_id, 100)
        self.assertEqual(record.reactor_username, "ivan")
        self.assertEqual(record.reactor_id, 42)
        self.assertEqual(record.stars_count, 7)

    def test_build_records_forward(self):
        client = FakeClient(
            entities={("user", 42): make_user(42, "ivan"),
                      ("channel", 555): make_channel(555, title="Source", username="src")},
        )
        forward = types.MessageFwdHeader(date=datetime(2024, 1, 1),
                                         from_id=types.PeerChannel(555), channel_post=777)
        document = types.Document(id=1, access_hash=1, dc_id=1, file_reference=b"",
                                  date=datetime(2024, 1, 1), mime_type="video/mp4", size=10,
                                  attributes=[types.DocumentAttributeVideo(1, 2, 2, False, False)])
        message = make_message(
            msg_id=200, media=types.MessageMediaDocument(document=document),
            fwd_from=forward,
            reactions=make_paid_reactions([types.MessageReactor(count=2,
                                                                peer_id=types.PeerUser(42))]),
        )
        records = self.run_async(
            build_records_for_message(client, message, "My Channel", self.cache))
        self.assertEqual(records[0].message_type, "video(forward)")
        self.assertEqual(records[0].original_channel, "src")   # username источника
        self.assertEqual(records[0].original_message_id, 777)
        self.assertEqual(records[0].current_message_id, 200)

    def test_no_stars_no_records(self):
        client = FakeClient()
        message = make_message(msg_id=1, text="hi", reactions=types.MessageReactions(
            results=[types.ReactionCount(reaction=types.ReactionEmoji("🔥"), count=5)]))
        self.assertEqual(self.run_async(
            build_records_for_message(client, message, "Channel", self.cache)), [])


# --------------------------------------------------------------------------- #
# Парсинг канала целиком
# --------------------------------------------------------------------------- #
class TestParseChannel(unittest.TestCase):
    def _messages(self):
        paid = make_paid_reactions([
            types.MessageReactor(count=3, peer_id=types.PeerUser(42)),
            types.MessageReactor(count=1, anonymous=True),
        ])
        return [
            make_message(msg_id=3, text="with stars", reactions=paid),
            types.MessageService(id=2, peer_id=types.PeerChannel(1),
                                             date=datetime(2024, 1, 1)),
            make_message(msg_id=1, text="no stars"),
        ]

    def test_parse_counts(self):
        client = FakeClient(messages=self._messages(),
                            entities={("user", 42): make_user(42, "ivan")})
        records: list[StarRecord] = []
        stats = asyncio.run(parse_channel(client, types.PeerChannel(1), 10, records,
                                          delay=(0, 0)))
        self.assertEqual(stats.scanned, 3)
        self.assertEqual(stats.skipped_service, 1)
        self.assertEqual(stats.with_stars, 1)
        self.assertEqual(stats.reactors, 2)
        self.assertEqual(len(records), 2)
        self.assertEqual([r.reactor_type for r in records], [REACTOR_USER, REACTOR_ANONYMOUS])

    def test_limit_respected(self):
        client = FakeClient(messages=self._messages())
        records: list[StarRecord] = []
        stats = asyncio.run(parse_channel(client, types.PeerChannel(1), 2, records,
                                          delay=(0, 0)))
        self.assertEqual(stats.scanned, 2)

    def test_flood_wait_continues(self):
        client = FakeClient(messages=self._messages(), flood_once=True,
                            entities={("user", 42): make_user(42, "ivan")})
        records: list[StarRecord] = []
        stats = asyncio.run(parse_channel(client, types.PeerChannel(1), 10, records,
                                          delay=(0, 0)))
        self.assertEqual(client.iter_calls, 2)      # была повторная попытка
        self.assertEqual(stats.scanned, 3)          # сообщение не обработано дважды
        self.assertEqual(len(records), 2)

    def test_channel_title(self):
        self.assertEqual(channel_title(make_channel(1, title="News", username="news")), "News")
        self.assertEqual(channel_title(make_channel(1, title="", username="news")), "news")


# --------------------------------------------------------------------------- #
# Экспорт
# --------------------------------------------------------------------------- #
def sample_records() -> list[StarRecord]:
    return [
        StarRecord("text", REACTOR_USER, "My Channel", 100, "My Channel", 100,
                   "ivan", 42, 7),
        StarRecord("video(forward)", REACTOR_ANONYMOUS, "My Channel", 99, "src", 777,
                   "", None, 1),
    ]


class TestExport(unittest.TestCase):
    def test_export_creates_xlsx(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = export_records(sample_records(), Path(tmp), "My|Channel")
            self.assertIsNotNone(path)
            assert path is not None
            self.assertTrue(path.exists())
            self.assertTrue(path.name.endswith("_My_Channel.xlsx"))

            workbook = load_workbook(path)
            sheet = workbook["Stars"]
            self.assertEqual([cell.value for cell in sheet[1]], list(COLUMNS))
            self.assertTrue(sheet["A1"].font.bold)
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertEqual(sheet.max_row, 3)

            # id остаются целыми числами (не float и не экспоненциальная запись)
            id_column = COLUMNS.index("current_message_id") + 1
            reactor_id_column = COLUMNS.index("reactor_id") + 1
            self.assertEqual(sheet.cell(row=2, column=id_column).value, 100)
            self.assertIsInstance(sheet.cell(row=2, column=id_column).value, int)
            self.assertEqual(sheet.cell(row=2, column=reactor_id_column).value, 42)
            self.assertIsInstance(sheet.cell(row=2, column=reactor_id_column).value, int)
            # анонимный отправитель — без id
            self.assertIsNone(sheet.cell(row=3, column=reactor_id_column).value)

    def test_no_records_no_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertIsNone(export_records([], Path(tmp), "Channel"))
            self.assertEqual(list(Path(tmp).iterdir()), [])

    def test_empty_file_on_demand(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = export_records([], Path(tmp), "Channel", create_empty_file=True)
            self.assertIsNotNone(path)
            sheet = load_workbook(path)["Stars"]  # type: ignore[arg-type]
            self.assertEqual([cell.value for cell in sheet[1]], list(COLUMNS))
            self.assertEqual(sheet.max_row, 1)

    def test_write_excel_directly(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "report.xlsx"
            write_excel(sample_records(), path)
            self.assertTrue(path.exists())

    def test_long_channel_name(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = export_records(sample_records(), Path(tmp), "C" * 500)
            assert path is not None
            self.assertLessEqual(len(path.name), 130)


class TestClientConnect(unittest.TestCase):
    """Тесты функции подключения client.connect()."""

    def test_connect_sync_is_connected(self):
        class MockClientSync:
            def __init__(self):
                self.connected = False
            async def connect(self):
                self.connected = True
            def is_connected(self) -> bool:
                return self.connected

        client = MockClientSync()
        asyncio.run(connect(client, attempts=1))
        self.assertTrue(client.connected)

    def test_connect_async_is_connected(self):
        class MockClientAsync:
            def __init__(self):
                self.connected = False
            async def connect(self):
                self.connected = True
            async def is_connected(self):
                return self.connected

        client = MockClientAsync()
        asyncio.run(connect(client, attempts=1))
        self.assertTrue(client.connected)

    def test_connect_failure_raises(self):
        class MockClientFail:
            async def connect(self):
                pass
            def is_connected(self) -> bool:
                return False

        with self.assertRaises(ConnectionFailureError):
            asyncio.run(connect(MockClientFail(), attempts=1))


if __name__ == "__main__":
    unittest.main(verbosity=2)
