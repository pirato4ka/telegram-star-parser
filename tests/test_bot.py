"""Тесты для Telegram-бота (bot package): конфиг, агрегация, форматирование, очередь, экспорт."""

import asyncio
import os
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pytest

from bot.aggregate import DonorAggregate, aggregate_star_records
from bot.config_bot import BotConfig, BotConfigError, load_bot_config
from bot.formatting import format_donors_table, format_queue_summary
from bot.queue_service import (
    ParseTask,
    QueueService,
    STATUS_CANCELLED,
    STATUS_COMPLETED,
    STATUS_PENDING,
    STATUS_RUNNING,
)
from exporter import DONOR_COLUMNS, export_donors_summary
from models import NOT_FOUND, REACTOR_ANONYMOUS, REACTOR_CHANNEL, REACTOR_USER, ParseStats, StarRecord
import pandas as pd
import openpyxl


def test_bot_config_validation():
    # Без токена
    if "BOT_TOKEN" in os.environ:
        del os.environ["BOT_TOKEN"]

    with tempfile.NamedTemporaryFile("w", delete=False) as f:
        f.write("[BOT]\nTOKEN=\n")
        cfg_path = f.name

    try:
        with pytest.raises(BotConfigError):
            load_bot_config(cfg_path)
    finally:
        os.remove(cfg_path)

    # С токеном и валидными полями
    with tempfile.NamedTemporaryFile("w", delete=False) as f:
        f.write(
            "[BOT]\n"
            "TOKEN=12345:ABC\n"
            "ALLOWED_USER_IDS=111, 222 , 333\n"
            "MAX_CONCURRENT_TASKS=2\n"
            "DONOR_THRESHOLD=100\n"
            "SEND_FILES=false\n"
            "DETAILED_EXCEL=true\n"
        )
        cfg_path = f.name

    try:
        cfg = load_bot_config(cfg_path)
        assert cfg.token == "12345:ABC"
        assert cfg.allowed_user_ids == {111, 222, 333}
        assert cfg.max_concurrent_tasks == 1  # форсируется в 1
        assert cfg.donor_threshold == 100
        assert cfg.send_files is False
        assert cfg.detailed_excel is True
    finally:
        os.remove(cfg_path)

    # Переменная окружения BOT_TOKEN переопределяет
    os.environ["BOT_TOKEN"] = "env_secret_token"
    try:
        with tempfile.NamedTemporaryFile("w", delete=False) as f:
            f.write("[BOT]\nTOKEN=file_token\n")
            cfg_path = f.name
        cfg = load_bot_config(cfg_path)
        assert cfg.token == "env_secret_token"
    finally:
        del os.environ["BOT_TOKEN"]
        os.remove(cfg_path)


def test_aggregate_records():
    records = [
        # alice: 2 поста, 100 + 50 звёзд = 150
        StarRecord("text", REACTOR_USER, "@ch1", 1, "@ch1", 1, "alice", 101, 100),
        StarRecord("text", REACTOR_USER, "@ch1", 2, "@ch1", 2, "alice", 101, 50),
        # bob: 1 пост, 200 звёзд
        StarRecord("text", REACTOR_USER, "@ch1", 1, "@ch1", 1, "bob", 102, 200),
        # not_found: 1 пост, 30 звёзд
        StarRecord("text", REACTOR_USER, "@ch1", 3, "@ch1", 3, NOT_FOUND, 103, 30),
        # anonymous: 2 поста, 20 + 10 звёзд = 30
        StarRecord("text", REACTOR_ANONYMOUS, "@ch1", 1, "@ch1", 1, "", None, 20),
        StarRecord("text", REACTOR_ANONYMOUS, "@ch2", 5, "@ch2", 5, "", None, 10),
    ]

    donors, anon = aggregate_star_records(records, threshold=80, only_above_threshold=False)

    # Порядок: bob (200), alice (150), id103 (not_found) (30)
    assert len(donors) == 3
    assert donors[0].reactor_username == "bob"
    assert donors[0].stars_total == 200
    assert donors[0].rank == 1

    assert donors[1].reactor_username == "alice"
    assert donors[1].stars_total == 150
    assert donors[1].posts_count == 2
    assert donors[1].rank == 2

    assert donors[2].reactor_username == "id103 (not_found)"
    assert donors[2].stars_total == 30
    assert donors[2].rank == 3

    assert anon is not None
    assert anon.reactor_username == "(анонимы)"
    assert anon.stars_total == 30
    assert anon.posts_count == 2
    assert set(anon.channels) == {"@ch1", "@ch2"}

    # Режим с фильтром по порогу (> 80)
    donors_thresh, anon_thresh = aggregate_star_records(records, threshold=80, only_above_threshold=True)
    assert len(donors_thresh) == 2
    assert [d.reactor_username for d in donors_thresh] == ["bob", "alice"]
    # Анонимы всегда остаются
    assert anon_thresh is not None
    assert anon_thresh.stars_total == 30


def test_export_donors_summary():
    donors = [
        DonorAggregate(rank=1, reactor_username="bob", reactor_id=102, reactor_type="user", stars_total=200, posts_count=1, channels=["@ch1"]),
        DonorAggregate(rank=2, reactor_username="alice", reactor_id=101, reactor_type="user", stars_total=150, posts_count=2, channels=["@ch1"]),
        DonorAggregate(rank=3, reactor_username="(анонимы)", reactor_id=None, reactor_type="anonymous", stars_total=30, posts_count=2, channels=["@ch1", "@ch2"]),
    ]
    stats = {
        "@ch1": ParseStats(requested=100, scanned=90, errors=1, unparsed=0, not_found=1),
    }

    with tempfile.TemporaryDirectory() as tmpdir:
        out_path = Path(tmpdir) / "report.xlsx"
        saved = export_donors_summary(
            donors=donors,
            channels_meta={"@ch1": {}},
            stats_by_channel=stats,
            requested_scope="Последние 100",
            output_path=str(out_path),
        )
        assert Path(saved).exists()

        wb = openpyxl.load_workbook(saved)
        assert "Donors" in wb.sheetnames
        assert "Summary" in wb.sheetnames

        ws_donors = wb["Donors"]
        headers = [cell.value for cell in ws_donors[1]]
        assert headers == list(DONOR_COLUMNS)

        ws_summary = wb["Summary"]
        summary_rows = {row[0]: row[1] for row in ws_summary.iter_rows(values_only=True) if row[0] is not None}
        assert summary_rows.get("Диапазон парсинга") == "Последние 100"
        assert summary_rows.get("Обработано сообщений (scanned)") == 90
        assert summary_rows.get("Результат неполный") == "Да"  # т.к. errors=1


def test_formatting_table():
    donors = [
        DonorAggregate(rank=1, reactor_username="user_<tag>", reactor_id=111, reactor_type="user", stars_total=100, posts_count=5, channels=["@ch"]),
    ]
    anon = DonorAggregate(rank=2, reactor_username="(анонимы)", reactor_id=None, reactor_type="anonymous", stars_total=20, posts_count=2, channels=["@ch"])

    stats = {
        "@ch": ParseStats(requested=50, scanned=50, errors=0, unparsed=0),
    }

    msgs = format_donors_table(donors, anon, ["@ch"], threshold=80, only_above_threshold=False, stats_by_channel=stats)
    assert len(msgs) >= 1
    # Проверяем экранирование HTML
    assert "&lt;tag&gt;" in msgs[0]
    assert "<tag>" not in msgs[0]
    assert "1 донатера, 120 звёзд (включая анонимов)" in msgs[0]
    assert "Обработано:" in msgs[0]


def test_queue_service_fifo_and_cancellation():
    async def _test():
        queue = QueueService()
        queue.start_worker()

        executed = []

        async def fake_task_coro(task, val):
            await asyncio.sleep(0.05)
            executed.append(val)

        task1 = ParseTask(
            task_id="t1",
            user_id=1,
            user_display="user1",
            task_type="parse",
            channels=["@ch1"],
            scope_desc="all",
            coro_func=fake_task_coro,
            coro_args=(1,),
        )
        task2 = ParseTask(
            task_id="t2",
            user_id=2,
            user_display="user2",
            task_type="parse",
            channels=["@ch2"],
            scope_desc="all",
            coro_func=fake_task_coro,
            coro_args=(2,),
        )

        pos1 = await queue.add_task(task1)
        pos2 = await queue.add_task(task2)
        assert pos1 == 1
        assert pos2 == 2

        # Отменяем task2 до того как он начался
        await queue.cancel_task(task2)
        assert task2.cancelled_by_user is True

        await asyncio.sleep(0.2)
        assert executed == [1]
        assert task1.status == STATUS_COMPLETED
        assert task2.status == STATUS_CANCELLED
        queue.stop_worker()

    asyncio.run(_test())


def test_year_range_filtering():
    """Проверка правильности фильтрации сообщений по годам в iter_message_batches."""
    from types import SimpleNamespace
    from handlers import iter_message_batches

    class DummyClient:
        def __init__(self, msgs):
            self.msgs = msgs

        async def iter_messages(self, entity, limit=None, offset_id=0, **kwargs):
            for m in self.msgs:
                if offset_id and m.id >= offset_id:
                    continue
                yield m

    m2024 = SimpleNamespace(id=1, date=datetime(2024, 6, 1, tzinfo=timezone.utc), reactions=None)
    m2023 = SimpleNamespace(id=2, date=datetime(2023, 6, 1, tzinfo=timezone.utc), reactions=None)
    m2022 = SimpleNamespace(id=3, date=datetime(2022, 6, 1, tzinfo=timezone.utc), reactions=None)
    m2021 = SimpleNamespace(id=4, date=datetime(2021, 6, 1, tzinfo=timezone.utc), reactions=None)
    m2020 = SimpleNamespace(id=5, date=datetime(2020, 6, 1, tzinfo=timezone.utc), reactions=None)

    client = DummyClient([m2024, m2023, m2022, m2021, m2020])

    async def _run():
        # Выбираем только 2023 и 2021 (пропуская 2024, 2022 и останавливаясь до 2020)
        s_date = datetime(2021, 1, 1, 0, 0, 0, tzinfo=timezone.utc)
        e_date = datetime(2023, 12, 31, 23, 59, 59, tzinfo=timezone.utc)
        selected_years = {2021, 2023}

        collected = []
        async for batch in iter_message_batches(
            client=client,
            entity="dummy",
            limit=100,
            batch_size=10,
            start_date=s_date,
            end_date=e_date,
            selected_years=selected_years,
        ):
            collected.extend(batch)

        assert [m.id for m in collected] == [2, 4]  # m2023 и m2021

    asyncio.run(_run())
