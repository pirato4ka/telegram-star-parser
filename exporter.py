"""Формирование Excel-отчёта и безопасное имя файла."""

from __future__ import annotations

import re
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import TYPE_CHECKING, Iterable, Optional, Sequence

import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from models import (
    COLUMNS,
    EXCEL_SHEET_NAME,
    ParseStats,
    StarRecord,
    USERNAME_MISSING,
)
from utils import (
    ILLEGAL_FILENAME_CHARS,
    WINDOWS_RESERVED_NAMES,
    clean_cell_value,
    collapse_spaces,
    get_logger,
    remove_invisible,
)

if TYPE_CHECKING:
    from bot.aggregate import DonorAggregate

# --- Лист «Донаты»: одна строка = один донат (пост + отправитель + звёзды) ---- #
# Поля 1:1 как на согласованном образце (скриншот 2): каждый донат отдельной
# строкой, со ссылкой на пост и количеством звёзд.
DONATION_COLUMNS = COLUMNS

DONATION_HEADERS = {name: name for name in DONATION_COLUMNS}

# --- Лист «Сводка»: агрегация по донатерам ------------------------------------ #
DONOR_COLUMNS = (
    "rank",
    "reactor_username",
    "reactor_id",
    "reactor_type",
    "stars_total",
    "posts_count",
    "channels",
)

DONOR_HEADERS = {
    "rank": "№",
    "reactor_username": "username",
    "reactor_id": "id",
    "reactor_type": "тип",
    "stars_total": "звёзд всего",
    "posts_count": "постов",
    "channels": "каналы",
}

DONATIONS_SHEET_NAME = "Донаты"
DONORS_SHEET_NAME = "Сводка"
INFO_SHEET_NAME = "Инфо"

logger = get_logger("exporter")

FALLBACK_CHANNEL_NAME = "channel"
DEFAULT_MAX_LENGTH = 100
MAX_SUFFIX_ATTEMPTS = 100

# Запрещённые в именах файлов символы (Windows) + управляющие.
ILLEGAL_CHARS_RE = re.compile(f"[{re.escape(ILLEGAL_FILENAME_CHARS)}\\u0000-\\u001f\\u007f-\\u009f]")

# Колонки, которые должны отображаться в Excel целыми числами (без экспоненциальной записи).
INTEGER_COLUMNS = ("current_message_id", "original_message_id", "reactor_id", "stars_count")

MIN_COLUMN_WIDTH = 10
MAX_COLUMN_WIDTH = 60


# --------------------------------------------------------------------------- #
# Имя файла
# --------------------------------------------------------------------------- #
def safe_filename(name: str, max_length: int = DEFAULT_MAX_LENGTH,
                  fallback: str = FALLBACK_CHANNEL_NAME) -> str:
    """Приводит название канала к безопасному имени файла для Windows.

    * удаляет невидимые символы и variation selectors (в т. ч. «скрытую» часть эмодзи);
    * заменяет `\\ / : * ? " < > |` и управляющие символы на `_`;
    * схлопывает пробелы, убирает пробелы/точки в конце;
    * обрезает до `max_length` символов (по границам символов, не по байтам);
    * для зарезервированных имён Windows (CON, NUL, COM1...) добавляет подчёркивание.
    """
    text = remove_invisible(name or "")
    text = ILLEGAL_CHARS_RE.sub("_", text)
    text = unicodedata.normalize("NFC", text)
    text = collapse_spaces(text)
    text = text.strip(" \t\r\n.")
    text = text.lstrip("-") if text.startswith("-") and len(text) > 1 else text

    if not text or not text.strip("_ ."):
        # Имя состоит только из мусорных символов (например "//" -> "__").
        return fallback

    if len(text) > max_length:
        text = text[:max_length].rstrip(" .")
        if not text:
            text = fallback

    # Имя не может заканчиваться точкой или пробелом (Windows).
    text = text.rstrip(" .")
    if not text:
        return fallback

    root = text.split(".")[0].upper()
    if root in WINDOWS_RESERVED_NAMES:
        text = f"_{text}"

    return text or fallback


def build_result_filename(channel_name: str, when: Optional[datetime] = None,
                          max_length: int = DEFAULT_MAX_LENGTH) -> str:
    """Формирует имя файла: `YYYY-MM-DD_HH-MM-SS_<канал>.xlsx`."""
    stamp = (when or datetime.now()).strftime("%Y-%m-%d_%H-%M-%S")
    channel = safe_filename(channel_name, max_length=max_length)
    if len(channel) > max_length:
        channel = channel[:max_length].rstrip(" .") or FALLBACK_CHANNEL_NAME
    return f"{stamp}_{channel}.xlsx"


def unique_path(directory: Path, filename: str,
                max_attempts: int = MAX_SUFFIX_ATTEMPTS) -> Path:
    """Возвращает свободный путь: `name.xlsx`, `name_1.xlsx`, `name_2.xlsx`, ..."""
    candidate = directory / filename
    if not candidate.exists():
        return candidate

    stem, suffix = Path(filename).stem, Path(filename).suffix
    for index in range(1, max_attempts + 1):
        candidate = directory / f"{stem}_{index}{suffix}"
        if not candidate.exists():
            return candidate
    # Крайний случай: добавляем метку времени.
    return directory / f"{stem}_{datetime.now().strftime('%H%M%S')}{suffix}"


# --------------------------------------------------------------------------- #
# Таблица
# --------------------------------------------------------------------------- #
def build_dataframe(records: Sequence[StarRecord]) -> pd.DataFrame:
    """Собирает DataFrame из записей (пустой — с правильными колонками).

    В колонке `reactor_username` — username донатера, а при его отсутствии
    «отсутствует» (`models.USERNAME_MISSING`).
    """
    rows = []
    for record in records:
        row = record.to_row()
        row[COLUMNS.index("reactor_username")] = record.username_for_report()
        rows.append([clean_cell_value(value) for value in row])
    frame = pd.DataFrame(rows, columns=list(COLUMNS), dtype=object)
    return frame


def _autosize(worksheet, frame: pd.DataFrame) -> None:
    """Ширина колонок по содержимому (с ограничениями)."""
    for index, column in enumerate(frame.columns, start=1):
        values = frame[column].tolist()
        longest = max([len(str(column))] + [len(str(v)) for v in values if v is not None] or [0])
        width = min(max(longest + 2, MIN_COLUMN_WIDTH), MAX_COLUMN_WIDTH)
        worksheet.column_dimensions[get_column_letter(index)].width = width


def _style_worksheet(worksheet, frame: pd.DataFrame) -> None:
    """Жирная шапка, закреплённая первая строка, числовой формат для id, автофильтр."""
    header_font = Font(bold=True)
    for cell in worksheet[1]:
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left", vertical="center")

    worksheet.freeze_panes = "A2"
    if len(frame.columns):
        last_column = get_column_letter(len(frame.columns))
        worksheet.auto_filter.ref = f"A1:{last_column}1"

    # id и количество звёзд — целыми числами, без экспоненциальной записи.
    for column_name in INTEGER_COLUMNS:
        if column_name not in frame.columns:
            continue
        index = list(frame.columns).index(column_name) + 1
        letter = get_column_letter(index)
        for cell in worksheet[letter][1:]:
            if cell.value is not None:
                cell.number_format = "0"

    # Ссылки на посты делаем кликабельными гиперссылками.
    if "post_link" in frame.columns:
        link_index = list(frame.columns).index("post_link") + 1
        link_letter = get_column_letter(link_index)
        link_font = Font(color="0563C1", underline="single")
        for cell in worksheet[link_letter][1:]:
            value = str(cell.value or "")
            if value.startswith(("https://", "http://")):
                cell.hyperlink = value
                cell.font = link_font

    _autosize(worksheet, frame)


def write_excel(records: Sequence[StarRecord], path: Path) -> None:
    """Записывает записи в .xlsx (лист `Stars`)."""
    frame = build_dataframe(records)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        frame.to_excel(writer, sheet_name=EXCEL_SHEET_NAME, index=False)
        _style_worksheet(writer.sheets[EXCEL_SHEET_NAME], frame)


def export_records(
    records: Iterable[StarRecord],
    output_dir: Path,
    channel_name: str,
    filename_max_length: int = DEFAULT_MAX_LENGTH,
    create_empty_file: bool = False,
) -> Optional[Path]:
    """Сохраняет отчёт и возвращает путь к файлу (или None, если сохранять нечего).

    При `PermissionError`/`OSError` (файл открыт в Excel) файл сохраняется
    с суффиксом `_1`, `_2` и т. д.
    """
    records = list(records)
    if not records and not create_empty_file:
        logger.info("Звёзд не найдено — файл не создаётся.")
        return None

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    filename = build_result_filename(channel_name, max_length=filename_max_length)
    path = unique_path(output_dir, filename)

    for attempt in range(MAX_SUFFIX_ATTEMPTS):
        try:
            write_excel(records, path)
            logger.info("Файл сохранён: %s (записей: %d)", path, len(records))
            return path
        except PermissionError as exc:
            logger.warning("Файл %s занят (%s) — пробуем другое имя.", path, exc)
        except OSError as exc:
            logger.warning("Ошибка записи %s (%s) — пробуем другое имя.", path, exc)

        stem, suffix = path.stem, path.suffix
        for index in range(1, MAX_SUFFIX_ATTEMPTS + 1):
            candidate = output_dir / f"{stem}_{index}{suffix}"
            if not candidate.exists():
                path = candidate
                break
        else:  # pragma: no cover - защитный вариант
            path = output_dir / f"{stem}_{datetime.now().strftime('%H%M%S')}{suffix}"
        _ = attempt

    logger.error("Не удалось сохранить файл после %d попыток.", MAX_SUFFIX_ATTEMPTS)
    return None


def build_donation_rows(records: Sequence[StarRecord]) -> list[dict]:
    """Строки листа «Донаты»: каждый донат отдельно, со ссылкой на пост.

    Состав и порядок полей — как на согласованном образце (скриншот 2):
    `message_type`, `reactor_type`, `current_channel`, `current_message_id`,
    `post_link`, `original_channel`, `original_message_id`, `reactor_username`,
    `reactor_id`, `stars_count`.

    Отличие ровно одно: в `reactor_username` пишется username донатера, а если
    его нет — «отсутствует» (вместо `not_found`/имени/пустой ячейки).
    """
    rows: list[dict] = []
    for record in records:
        row = record.to_dict()
        row["reactor_username"] = record.username_for_report()
        rows.append(row)
    return rows


def export_donors_summary(
    donors: list[DonorAggregate],
    channels_meta: dict,
    stats_by_channel: dict[str, ParseStats],
    requested_scope: str,
    output_path: str,
    records: Optional[Sequence[StarRecord]] = None,
) -> str:
    """Формирует Excel-отчёт бота.

    Листы:
    * «Донаты» — каждая отправка звёзд отдельной строкой: поля 1:1 как на
      согласованном образце (`message_type`, `reactor_type`, `current_channel`,
      `current_message_id`, `post_link`, `original_channel`, `original_message_id`,
      `reactor_username`, `reactor_id`, `stars_count`). В `reactor_username`
      пишется username донатера, а если его нет — «отсутствует»;
    * «Сводка» — агрегация по донатерам (кто сколько всего отправил);
    * «Инфо» — параметры парсинга и полнота результата.

    Возвращает абсолютный путь к сохранённому файлу.
    """
    path = Path(output_path).resolve()
    path.parent.mkdir(parents=True, exist_ok=True)

    # 1. Лист «Донаты» — построчно, без суммирования
    donation_rows = build_donation_rows(records or [])
    df_donations = pd.DataFrame(
        [[clean_cell_value(row[name]) for name in DONATION_COLUMNS] for row in donation_rows],
        columns=list(DONATION_COLUMNS),
    )
    df_donations = df_donations.rename(columns=DONATION_HEADERS)

    # 2. Лист «Сводка» — агрегация по донатерам
    donor_rows = []
    for d in donors:
        channels_str = ", ".join(d.channels) if isinstance(d.channels, (list, set, tuple)) else str(d.channels or "")
        donor_rows.append({
            "rank": d.rank,
            "reactor_username": clean_cell_value(d.reactor_username) or USERNAME_MISSING,
            "reactor_id": d.reactor_id if d.reactor_id is not None else "",
            "reactor_type": d.reactor_type,
            "stars_total": d.stars_total,
            "posts_count": d.posts_count,
            "channels": channels_str,
        })

    df_donors = pd.DataFrame(donor_rows, columns=list(DONOR_COLUMNS))
    df_donors = df_donors.rename(columns=DONOR_HEADERS)

    # 3. Лист «Инфо» — параметры парсинга
    requested_values = [s.requested for s in stats_by_channel.values()] if stats_by_channel else []
    total_requested = sum(v for v in requested_values if v) if any(v for v in requested_values) else 0
    total_scanned = sum(s.scanned for s in stats_by_channel.values()) if stats_by_channel else 0
    total_errors = sum(s.errors for s in stats_by_channel.values()) if stats_by_channel else 0
    total_unparsed = sum(s.unparsed for s in stats_by_channel.values()) if stats_by_channel else 0
    total_not_found = sum(s.not_found for s in stats_by_channel.values()) if stats_by_channel else 0
    total_anon = sum(1 for d in donors if d.reactor_type == "anonymous")
    total_stars = sum(d.stars_total for d in donors)
    total_donations = len(donation_rows)

    is_incomplete = (total_errors + total_unparsed) > 0 or any(s.interrupted for s in stats_by_channel.values())

    channel_names_list = list(stats_by_channel.keys())
    if not channel_names_list and channels_meta:
        channel_names_list = list(channels_meta.keys())

    processed_value = str(total_scanned)
    if total_requested and total_requested != total_scanned:
        processed_value = f"{total_scanned} из {total_requested}"

    summary_data = [
        {"Параметр": "Список каналов", "Значение": ", ".join(channel_names_list)},
        {"Параметр": "Диапазон парсинга", "Значение": requested_scope},
        {"Параметр": "Дата формирования", "Значение": datetime.now().strftime("%Y-%m-%d %H:%M:%S")},
        {"Параметр": "Записей о донатах", "Значение": total_donations},
        {"Параметр": "Всего звёзд", "Значение": total_stars},
        {"Параметр": "Обработано постов", "Значение": processed_value},
        {"Параметр": "Ошибок и непрочитанных (errors + unparsed)", "Значение": total_errors + total_unparsed},
        {"Параметр": "Не расшифровано (not_found)", "Значение": total_not_found},
        {"Параметр": "Анонимных записей (anonymous)", "Значение": total_anon},
        {"Параметр": "Результат неполный", "Значение": "Да" if is_incomplete else "Нет"},
    ]
    df_summary = pd.DataFrame(summary_data)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df_donations.to_excel(writer, sheet_name=DONATIONS_SHEET_NAME, index=False)
        _style_donations_worksheet(writer.sheets[DONATIONS_SHEET_NAME], df_donations)

        df_donors.to_excel(writer, sheet_name=DONORS_SHEET_NAME, index=False)
        _style_donors_worksheet(writer.sheets[DONORS_SHEET_NAME], df_donors)

        df_summary.to_excel(writer, sheet_name=INFO_SHEET_NAME, index=False)
        _style_summary_worksheet(writer.sheets[INFO_SHEET_NAME], df_summary)

    logger.info("Excel-отчёт бота сохранён: %s (донатов: %d, донатеров: %d)",
                path, total_donations, len(donors))
    return str(path)


def _style_donations_worksheet(worksheet, frame: pd.DataFrame) -> None:
    """Оформление листа «Донаты»: шапка, фильтр, кликабельные ссылки, целые числа."""
    header_font = Font(bold=True)
    for cell in worksheet[1]:
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left", vertical="center")

    worksheet.freeze_panes = "A2"
    if len(frame.columns):
        last_column = get_column_letter(len(frame.columns))
        worksheet.auto_filter.ref = f"A1:{last_column}1"

    headers = list(frame.columns)
    for column_name in INTEGER_COLUMNS:
        if column_name not in headers:
            continue
        letter = get_column_letter(headers.index(column_name) + 1)
        for cell in worksheet[letter][1:]:
            if cell.value is None or str(cell.value).strip() == "":
                continue
            try:
                cell.value = int(cell.value)
                cell.number_format = "0"
            except (ValueError, TypeError):
                pass

    if "post_link" in headers:
        link_letter = get_column_letter(headers.index("post_link") + 1)
        link_font = Font(color="0563C1", underline="single")
        for cell in worksheet[link_letter][1:]:
            value = str(cell.value or "")
            if value.startswith(("https://", "http://")):
                cell.hyperlink = value
                cell.font = link_font

    _autosize(worksheet, frame)


def _style_donors_worksheet(worksheet, frame: pd.DataFrame) -> None:
    """Оформление листа «Сводка»."""
    header_font = Font(bold=True)
    for cell in worksheet[1]:
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left", vertical="center")

    worksheet.freeze_panes = "A2"
    if len(frame.columns):
        last_column = get_column_letter(len(frame.columns))
        worksheet.auto_filter.ref = f"A1:{last_column}1"

    # №, id, звёзд всего, постов — целыми числами
    int_cols = tuple(DONOR_HEADERS[key] for key in ("rank", "reactor_id", "stars_total", "posts_count"))
    for col_name in int_cols:
        if col_name not in frame.columns:
            continue
        index = list(frame.columns).index(col_name) + 1
        letter = get_column_letter(index)
        for cell in worksheet[letter][1:]:
            if cell.value is not None and str(cell.value).strip() != "":
                try:
                    cell.value = int(cell.value)
                    cell.number_format = "0"
                except (ValueError, TypeError):
                    pass

    _autosize(worksheet, frame)


def _style_summary_worksheet(worksheet, frame: pd.DataFrame) -> None:
    """Оформление листа «Инфо»."""
    header_font = Font(bold=True)
    for cell in worksheet[1]:
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left", vertical="center")

    worksheet.freeze_panes = "A2"
    _autosize(worksheet, frame)


__all__ = [
    "DONATION_COLUMNS", "DONATION_HEADERS", "DONATIONS_SHEET_NAME", "DONOR_COLUMNS",
    "DONOR_HEADERS", "DONORS_SHEET_NAME", "INFO_SHEET_NAME",
    "build_dataframe", "build_donation_rows", "build_result_filename",
    "export_donors_summary", "export_records", "safe_filename", "unique_path",
    "write_excel",
]
